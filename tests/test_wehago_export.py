from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from core.wehago_export import (
    ExportFormatError, parse_export_filename, read_wehago_export,
)

H1 = ["사원코드", "사원명", "부서", "직급", "직종", "수당", None, None, "공제", None, None, "차인지급액"]
H2 = [None, None, None, None, None, "기본급", "식대", "지급액계", "국민연금", "소득세", "공제액계", None]


def build(tmp_path, rows, total=None, name="가나다센터 - 202609.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(H1)
    ws.append(H2)
    for row in rows:
        ws.append(row)
    if total is not None:
        ws.append(total)
    path = tmp_path / name
    wb.save(path)
    return path


def emp(code, name, base, meal, pension, tax):
    gross, ded = base + meal, pension + tax
    return [code, name, "", "직급없음", "", base, meal, gross, pension, tax, ded, gross - ded]


def total_of(*rows):
    sums = [sum(r[i] for r in rows) for i in range(5, 12)]
    return ["합계", None, None, None, None] + [f"{v:,}" for v in sums]


def test_reads_codes_as_text_and_amounts(tmp_path):
    a = emp("1", "가", 1000000, 200000, 45000, 10000)
    b = emp("01", "나", 2000000, 0, 90000, 30000)
    export = read_wehago_export(build(tmp_path, [a, b], total_of(a, b)))
    assert [p.employee_code for p in export.people] == ["1", "01"]
    assert export.people[0].amounts.net == Decimal(1145000)
    assert export.totals.net == Decimal(1145000 + 1880000)


def test_duplicate_code_fails(tmp_path):
    a = emp("1", "가", 100, 0, 0, 0)
    with pytest.raises(ExportFormatError, match="반복"):
        read_wehago_export(build(tmp_path, [a, a], total_of(a, a)))


def test_item_columns_are_not_authoritative(tmp_path):
    # Observed in real exports: WEHAGO 지급액계 can be below the item sum.
    a = emp("1", "가", 100, 50, 0, 0)
    a[7] = 120
    a[11] = 120
    export = read_wehago_export(build(tmp_path, [a], total_of(a)))
    assert export.people[0].amounts.gross == Decimal(120)


def test_row_arithmetic_mismatch_fails(tmp_path):
    a = emp("1", "가", 100, 0, 0, 0)
    a[11] = 99
    with pytest.raises(ExportFormatError, match="차인지급액"):
        read_wehago_export(build(tmp_path, [a], total_of(a)))


def test_total_row_mismatch_and_missing_total_fail(tmp_path):
    a = emp("1", "가", 100, 0, 0, 0)
    bad = total_of(a)
    bad[-1] = "1"
    with pytest.raises(ExportFormatError, match="합계"):
        read_wehago_export(build(tmp_path, [a], bad))
    with pytest.raises(ExportFormatError, match="합계 행"):
        read_wehago_export(build(tmp_path, [a], None, name="x - 202609.xlsx"))


def test_not_an_export_fails(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["성명", "기본급"])
    path = tmp_path / "x.xlsx"
    wb.save(path)
    with pytest.raises(ExportFormatError):
        read_wehago_export(path)


def test_filename_identity():
    ident = parse_export_filename("업체079 - 202608 (1).xlsx")
    assert (ident.company, ident.year, ident.month) == ("업체079", 2026, 8)
    with pytest.raises(ExportFormatError):
        parse_export_filename("260921_업체04008(급여).xlsx")


REFERENCE_DIR = Path(__file__).resolve().parents[1] / "data" / "reference_files"


def _raw_exports():
    if not REFERENCE_DIR.is_dir():
        return []
    found = []
    for path in sorted(REFERENCE_DIR.glob("*.xlsx")):
        wb = openpyxl.load_workbook(path, read_only=True)
        first = next(wb.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True))
        wb.close()
        if tuple(first[:3]) == ("사원코드", "사원명", "부서") and "수당" in first:
            found.append(path)
    return found


def test_all_stored_raw_wehago_exports_parse():
    paths = _raw_exports()
    if not paths:
        pytest.skip("저장된 위하고 원본 내보내기 파일 없음")
    failures = {}
    for path in paths:
        try:
            read_wehago_export(path)
        except ExportFormatError as exc:
            failures[path.name] = str(exc)
    assert not failures, failures
