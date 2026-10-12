"""Verify mapping inference and duplicate statutory deduction protection."""
from pathlib import Path
from core.canon import norm
from core.engine import process
from core.export import export_wehago
from core.mapping_learn import collect_contexts, learn_from_contexts
from core.reconcile import learned_mapping_key

FIXTURES = Path(__file__).parent / 'fixtures' / 'golden'
SRC = FIXTURES / 'sample3-source.xlsx'
REF = FIXTURES / 'sample3-wehago-before.xlsx'

def _row(result, label):
    return next(r for r in result.mapping_rows if r.kind != '미매칭' and r.item_label == label)

def test_learns_mapping_that_a_person_changed(tmp_path):
    base = process('sample3', SRC, REF)
    row = _row(base, '기본급')
    other = next(c.key for c in base.profile.allowance_columns if c.key != row.dest_key)
    answer = process('sample3', SRC, REF, manual_overrides={row.source_col: other})
    answer_path = tmp_path / 'answer.xlsx'
    export_wehago(answer.profile, answer.records, answer_path)
    _, contexts = collect_contexts(lambda: process('sample3', SRC, answer_path))
    learned = learn_from_contexts(contexts)
    key = norm(learned_mapping_key(row.group, row.item_label))
    assert learned is not None
    assert learned.assignment[key] == other
    assert learned.mismatches_after < learned.mismatches_before
    assert learned.mismatches_after == 0
    again = process('sample3', SRC, REF, learned_mappings=learned.changed)
    assert _row(again, '기본급').dest_key == other

def test_nothing_changes_when_engine_already_matches_answer(tmp_path):
    base = process('sample3', SRC, REF)
    answer_path = tmp_path / 'answer.xlsx'
    export_wehago(base.profile, base.records, answer_path)
    _, contexts = collect_contexts(lambda: process('sample3', SRC, answer_path))
    learned = learn_from_contexts(contexts)
    assert learned is not None
    assert learned.mismatches_before == 0
    assert learned.changed == {}
    base_row = _row(base, "기본급")
    assert norm(learned_mapping_key(base_row.group, base_row.item_label)) in learned.verified
    row = _row(base, '기본급')
    assert norm(learned_mapping_key(row.group, row.item_label)) in learned.verified

def test_learn_sink_is_off_outside_collect():
    import core.reconcile as rec
    assert rec._LEARN_SINK is None
    process('sample3', SRC, REF)
    assert rec._LEARN_SINK is None

def test_learned_pension_mapping_does_not_restore_duplicate_column(tmp_path):
    import openpyxl
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    import test_source_template as T
    src, ref = tmp_path/'source.xlsx', tmp_path/'reference.xlsx'
    T._original(src, 8, T.PEOPLE)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    ws.cell(3, 8, '국민연금')
    for row in range(4, 7):
        ws.cell(row, 8, ws.cell(row, 5).value)
    wb.save(src)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['사원코드','사원명','기본급','지급액계','국민연금','소득세','공제액계','차인지급액'])
    for i, (name, base, _, pension, tax) in enumerate(T.PEOPLE, 1):
        ws.append([i,name,base,base,pension,tax,pension+tax,base-pension-tax])
    wb.save(ref)
    result = process('test', src, ref, learned_mappings={norm(learned_mapping_key('공제','국민연금')):'국민연금'})
    assert result is not None
    expected = {name:pension for name,_,_,pension,_ in T.PEOPLE}
    for record in result.records:
        assert record.deductions['국민연금'] == expected[record.사원명]
