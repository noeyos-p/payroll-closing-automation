"""End-to-end regression against real, already-macro-processed workbooks.

Each fixture in tests/fixtures/golden/ was split (see
/private/tmp/.../make_golden_fixtures.py, run once against 파일/*.xlsx — the
originals are untouched) out of a real client's June 2026 payroll workbook
that already contained '원본'(or 'Sheet2')/'Sheet1'/'정리본'/'검증결과'/
'항목매핑결과' sheets — i.e. a real run of
파일/이지케어_Sheet1_현재시트_수정본_v3.bas, done by hand in Excel. The
'-expected.xlsx' 검증결과 sheet is that real run's own answer key.

`SAMPLE4_KNOWN_DIVERGENCES` documents the one confirmed case where this
port's output is *more* correct than the golden answer key, not less: the
real macro's fixed row constants (SRC_H1=3/SRC_H2=5) misread this
particular client's 2-row header (rows 3-4, data at row 5) and silently
dropped the "연장근로" column for everyone — invisible for every employee
except the two who actually had a nonzero value there that month. This
port's core/sheet_layout.py detects the header/data rows from the sheet's
own structure instead of a hardcoded constant, so it maps that column
correctly. Confirmed by hand: 사원023's and 사원024's 연장근로 amounts in
their source row account for the entire gap against the golden 지급액계/
차인지급액 (203,820원 and 136,740원)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from core.engine import process

FIXTURES = Path(__file__).parent / "fixtures" / "golden"

GOLDEN_CASES = ["sample1", "sample3", "sample2", "sample4", "sample5"]

SAMPLE4_KNOWN_DIVERGENCES = {"사원023", "사원024"}

# 2026-10-07: sample1(업체075)의 원본 "식대" 열과 "기타공제 > 본인부담금"
# 열은 옛 매크로가 못 붙였던 항목이다(위하고 열 이름이 "식대(비)"/"본인부담금공제"라
# 글자가 달라서). 사용자가 이번 달 실제로 손으로 맞춰 올린 파일(작업파일\Exel의
# 업체030(부천)/의정부통합돌봄 9월지급분)은 둘 다 그 열에 금액을 넣었으므로, 이제
# 앱도 그렇게 매칭한다. 정답키(매크로 출력)와의 차이는 정확히 이 두 원본 열의
# 금액만큼이어야 한다 — 그 금액을 원본에서 직접 읽어 기대값에 더해서 비교한다.
SAMPLE1_NEWLY_MAPPED = {"식대": "pay", "본인부담금": "ded"}


def _sample1_adjustments() -> dict[str, tuple[float, float]]:
    from core.engine import _source_candidates
    from core.sheet_layout import source_item

    ws, layout = _source_candidates(FIXTURES / "sample1-source.xlsx")[0]
    cols = {
        c: SAMPLE1_NEWLY_MAPPED[source_item(ws, layout, c).strip()]
        for c in range(1, ws.max_column + 1)
        if source_item(ws, layout, c).strip() in SAMPLE1_NEWLY_MAPPED
    }
    assert set(cols.values()) == {"pay", "ded"}
    adjustments: dict[str, tuple[float, float]] = {}
    for r in range(layout.data_start_row, ws.max_row + 1):
        name = ws.cell(row=r, column=layout.name_col).value
        if not isinstance(name, str) or not name.strip():
            continue
        pay = ded = 0.0
        for c, kind in cols.items():
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)):
                if kind == "pay":
                    pay += v
                else:
                    ded += v
        adjustments[name.strip()] = (pay, ded)
    return adjustments


def _load_expected_verification(short: str) -> dict[str, dict]:
    workbook = openpyxl.load_workbook(FIXTURES / f"{short}-expected.xlsx", data_only=True)
    worksheet = workbook["검증결과"]
    expected: dict[str, dict] = {}
    for r in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=r, column=1).value
        if not name:
            continue
        expected[name] = {
            "status": worksheet.cell(row=r, column=2).value,
            "dst_pay": worksheet.cell(row=r, column=4).value or 0,
            "dst_ded": worksheet.cell(row=r, column=7).value or 0,
            "dst_net": worksheet.cell(row=r, column=10).value or 0,
        }
    return expected


@pytest.mark.parametrize("short", GOLDEN_CASES)
def test_reconcile_matches_golden_verification(short: str) -> None:
    result = process(
        short,
        FIXTURES / f"{short}-source.xlsx",
        FIXTURES / f"{short}-wehago-before.xlsx",
    )
    assert result is not None

    expected = _load_expected_verification(short)
    got = {row.name: row for row in result.verification_rows}

    assert set(got) == set(expected), (short, set(expected) - set(got), set(got) - set(expected))

    adjustments = _sample1_adjustments() if short == "sample1" else {}
    mismatches = []
    for name, exp in expected.items():
        if short == "sample4" and name in SAMPLE4_KNOWN_DIVERGENCES:
            continue
        add_pay, add_ded = adjustments.get(name, (0.0, 0.0))
        if exp["status"] in ("정상", "차이", "추가됨", "추가됨_차이"):
            exp = dict(exp)
            exp["dst_pay"] += add_pay
            exp["dst_ded"] += add_ded
            exp["dst_net"] += add_pay - add_ded
        row = got[name]
        if (
            abs((row.dst_pay or 0) - exp["dst_pay"]) > 1
            or abs((row.dst_ded or 0) - exp["dst_ded"]) > 1
            or abs((row.dst_net or 0) - exp["dst_net"]) > 1
        ):
            mismatches.append((name, exp, row))
    assert not mismatches, mismatches


@pytest.mark.parametrize("short", GOLDEN_CASES)
def test_reconcile_records_match_정리본_지급액계(short: str) -> None:
    result = process(
        short,
        FIXTURES / f"{short}-source.xlsx",
        FIXTURES / f"{short}-wehago-before.xlsx",
    )
    assert result is not None

    workbook = openpyxl.load_workbook(FIXTURES / f"{short}-expected.xlsx", data_only=True)
    worksheet = workbook["정리본"]
    header = [worksheet.cell(row=1, column=c).value for c in range(1, worksheet.max_column + 1)]
    pay_col = header.index("지급액계") + 1
    code_col = header.index("사원코드") + 1

    name_col = header.index("사원명") + 1
    expected_pay_by_code: dict[str, float] = {}
    expected_pay_by_name: dict[str, float] = {}
    for r in range(2, worksheet.max_row + 1):
        code = worksheet.cell(row=r, column=code_col).value
        if code is None or str(code).strip() == "":
            continue
        pay = worksheet.cell(row=r, column=pay_col).value or 0
        # "사원코드 없음" (newly-added employee) isn't a unique key — several
        # rows can share that literal text — so those fall back to name.
        if str(code) == "사원코드 없음":
            expected_pay_by_name[str(worksheet.cell(row=r, column=name_col).value)] = pay
        else:
            expected_pay_by_code[str(code)] = pay

    adjustments = _sample1_adjustments() if short == "sample1" else {}
    mismatches = []
    for record in result.records:
        if record.사원코드 == "사원코드 없음":
            expected_pay = expected_pay_by_name.get(record.사원명)
        else:
            expected_pay = expected_pay_by_code.get(record.사원코드)
        if expected_pay is None:
            continue  # a newly-added row this run has no golden counterpart
        if short == "sample4" and record.사원명 in SAMPLE4_KNOWN_DIVERGENCES:
            continue
        expected_pay += adjustments.get(record.사원명, (0.0, 0.0))[0]
        if abs(record.지급액계 - expected_pay) > 1:
            mismatches.append((record.사원코드, record.사원명, record.지급액계, expected_pay))
    assert not mismatches, mismatches


@pytest.mark.parametrize("short", GOLDEN_CASES)
def test_reconcile_grand_total_matches_source_ledger(short: str) -> None:
    """core/reconcile.py's `grand_total_*` — a whole-run check against the
    source ledger's OWN printed 합계 row, independent of any one employee's
    matching. Only meaningful for a fixture whose source sheet actually has
    a detectable 합계 row (sample3/sample5 do; the others don't, and
    `grand_total_source` is correctly empty for those — nothing to assert
    beyond "no false mismatch")."""
    result = process(
        short,
        FIXTURES / f"{short}-source.xlsx",
        FIXTURES / f"{short}-wehago-before.xlsx",
    )
    assert result is not None
    assert not result.grand_total_mismatches, (short, result.grand_total_source, result.grand_total_computed)
