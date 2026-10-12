import asyncio
from dataclasses import replace
from datetime import date
from decimal import Decimal

import openpyxl

from core.payroll_closing_run import discover_jobs, outcome_rows, run_folder
from core.payroll_source_check import SourceNet
from core.payroll_workflow import Employee, PayrollSnapshot

HEADER = ["번호", "업체명", "저장폴더", "사업자번호", "귀속지급구분", "사업소득여부", "원천세신고여부",
          "급여입력여부", "신고서저장\n완료여부", "납부구분", "국세\n납부서", "지방세\n납부서", "메일", "비고"]


def company_list(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "업체"
    ws.append(HEADER)
    for i, (name, number, h) in enumerate(rows, start=1):
        ws.append([i, name, "마음손", number, "차월", "X", "O", h, None, None, None, None, None, None])
    path = tmp_path / "업체리스트(10월지급).xlsx"
    wb.save(path)
    return path


def source_folder(tmp_path, names):
    folder = tmp_path / "src"
    folder.mkdir()
    for name in names:
        (folder / name).write_bytes(b"")
    return folder


NETS = {"가나다센터 9월귀속 10월지급 급여대장(급여).xlsx": [("홍길동", 100), ("김철수", 50)],
        "가나다센터 9월귀속 10월지급 급여대장(추급).xlsx": [("홍길동", 30)],
        "라마바요양원 9월귀속 10월지급 급여대장.xlsx": [("사원015", 70)],
        "이미완료센터 9월귀속 10월지급 급여대장.xlsx": [("사원016", 10)],
        "없는업체 9월귀속 10월지급 급여대장.xlsx": [("최", 1)]}


def read_net(path):
    people = [SourceNet(n, Decimal(v), str(path), 7) for n, v in NETS[path.name]]
    return people, sum((p.net for p in people), Decimal(0))


class Adapter:
    """Empty payments; per (company, kind) the stored nets can be overridden."""
    def __init__(self, stored_override=None):
        self.stored_override = stored_override or {}
        self.snapshots = {}
        self.uploads = []

    def _key(self, payment):
        return payment.company_id, payment.payment_kind

    async def list_closed(self, payment):
        return False

    async def read_payroll(self, payment):
        return self.snapshots.get(self._key(payment), PayrollSnapshot(payment, (), None, False))

    async def list_employees(self, payment):
        names = {"홍길동": "001", "김철수": "002", "사원015": "003", "사원016": "004"}
        return tuple(Employee(code, name) for name, code in names.items())

    async def register_temporary(self, payment, request):
        raise AssertionError("not expected")

    async def copy_previous_payment(self, payment):
        return replace(payment, payment_date=date(2026, 10, 25))

    async def upload_to_created(self, original):
        self.uploads.append(original.payment)
        people = original.people
        override = self.stored_override.get(self._key(original.payment))
        if override:
            people = tuple(replace(p, amounts=replace(p.amounts, net=Decimal(override)))
                           for p in people)
        totals = replace(original.totals, net=sum((p.amounts.net for p in people), Decimal(0)))
        self.snapshots[self._key(original.payment)] = PayrollSnapshot(
            original.payment, people, totals, True, revision="r1")
        return True

    async def close_if_unchanged(self, snapshot):
        key = self._key(snapshot.payment)
        self.snapshots[key] = replace(self.snapshots[key], closed=True)
        return True


def test_discover_jobs_reads_company_kind_and_period(tmp_path):
    folder = source_folder(tmp_path, list(NETS) + ["~$임시.xlsx", "다른달 8월귀속 9월지급 급여대장.xlsx",
                                                   "둘다 9월귀속 10월지급 급여대장(급여)(급+상).xlsx"])
    jobs, problems = discover_jobs(folder, 2026, 10)
    assert {(j.company_name, j.kind, j.accrual_month, j.kind_assumed) for j in jobs} >= {
        ("가나다센터", "급여", 9, False), ("가나다센터", "추급", 9, False), ("라마바요양원", "급여", 9, True)}
    assert {p.status for p in problems} == {"unmatched", "blocked"}
    assert all(j.accrual_year == 2026 for j in jobs)


def test_mixed_kinds_write_ox_and_skip_completed(tmp_path):
    list_path = company_list(tmp_path, [("가나다센터", "1234567890", None),
                                        ("라마바요양원", "2234567890", None),
                                        ("이미완료센터", "3234567890", "O")])
    folder = source_folder(tmp_path, list(NETS))
    adapter = Adapter({("1234567890", "추급"): 29})  # 추급 stored 1원 short → blocked
    outcomes, update = asyncio.run(run_folder(folder, list_path, 2026, 10, adapter,
                                              read_net=read_net, write_list=True))
    status = {(o.job.company_name if o.job else o.path.name, o.job.kind if o.job else ""): o.status
              for o in outcomes}
    assert status[("가나다센터", "급여")] == "closed"
    assert status[("가나다센터", "추급")] == "blocked"
    assert status[("라마바요양원", "급여")] == "closed"
    assert status[("이미완료센터", "급여")] == "skipped"
    assert status[("없는업체", "급여")] == "unmatched"
    ws = openpyxl.load_workbook(list_path)["업체"]
    assert ws["H2"].value == "O,X" and "추급: X" in ws["N2"].value
    # The list note carries only reasons for unclosed kinds; informational
    # notes (kind assumed from a kindless file name) go to the result table.
    assert ws["H3"].value == "O" and ws["N3"].value == "[급여마감] 급여: O"
    assert any(r["업체명"] == "라마바요양원" and "급여로 처리" in r["사유"]
               for r in outcome_rows(outcomes))
    assert ws["H4"].value == "O" and ws["N4"].value is None
    assert update.backup_path is not None
    assert {p.company_id for p in adapter.uploads} == {"1234567890", "2234567890"}
    rows = outcome_rows(outcomes)
    assert any(r["결과"] == "보류" and "원본" in r["사유"] for r in rows)


def test_dry_run_leaves_list_untouched(tmp_path):
    list_path = company_list(tmp_path, [("라마바요양원", "2234567890", None)])
    before = list_path.read_bytes()
    folder = source_folder(tmp_path, ["라마바요양원 9월귀속 10월지급 급여대장.xlsx"])
    outcomes, update = asyncio.run(run_folder(folder, list_path, 2026, 10, Adapter(),
                                              read_net=read_net, write_list=False))
    assert update is None and list_path.read_bytes() == before
    assert outcomes[0].status == "closed"


def test_all_blocked_company_keeps_h_blank(tmp_path):
    list_path = company_list(tmp_path, [("라마바요양원", "2234567890", None)])
    folder = source_folder(tmp_path, ["라마바요양원 9월귀속 10월지급 급여대장.xlsx"])
    adapter = Adapter({("2234567890", "급여"): 1})
    asyncio.run(run_folder(folder, list_path, 2026, 10, adapter, read_net=read_net, write_list=True))
    ws = openpyxl.load_workbook(list_path)["업체"]
    assert ws["H2"].value is None and "급여: X" in ws["N2"].value
