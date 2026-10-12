"""Regression tests for the 2026-10-07 generalization round (9월지급 급여대장
전체를 사용자가 손으로 맞춘 업로드 파일과 대조하며 고친 것들). Every sheet
here is built in memory — no external sample folder needed."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from core.canon import canon
from core.engine import process
from core.sheet_layout import find_total_row


def _save(rows: list[list[object]], path: Path, merges: list[str] = ()) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    for m in merges:
        ws.merge_cells(m)
    wb.save(path)
    return path


def _dest(path: Path, items: list[str], deductions: list[str], people: list[tuple[str, str]]) -> Path:
    """이 앱이 내보내는 것과 같은 한 줄 머리글 위하고 양식."""
    header = ["사원코드", "사원명", *items, "지급액계", *deductions, "공제액계", "차인지급액"]
    rows = [header]
    for code, name in people:
        rows.append([code, name] + [0] * (len(header) - 2))
    return _save(rows, path)


@pytest.mark.parametrize(
    "text, expected",
    [
        ("식대(비)", "식대"),
        ("(비)식대", "식대"),
        ("연장수당(과)", "연장근로수당"),
        ("고용보험 (0.9/100)", "고용보험"),
        ("기 본 급\n①", "기본급"),
        ("급여합계\n\n(①+②+③)", "지급액계"),
        ("장기근속공제", "장기근속수당공제"),
        ("실업급여", "고용보험"),
        ("본인부담", "본인부담"),  # 접미사 제거로 빈 이름이 되면 안 됨
    ],
)
def test_canon_variants(text: str, expected: str) -> None:
    assert canon(text) == expected


def test_total_row_with_headcount_suffix(tmp_path: Path) -> None:
    path = _save([["순번", "성명"], [1, "가"], ["합   계 (총  1명)", None]], tmp_path / "t.xlsx")
    ws = openpyxl.load_workbook(path).active
    assert find_total_row(ws, 2, 3) == 3


def test_unknown_dest_item_classified_by_position(tmp_path: Path) -> None:
    """위하고 열 "처우개선비"는 사전에 없지만 지급액계 앞에 있으니 수당이다 —
    원본의 같은 이름 열이 그대로 매칭돼야 한다."""
    src = _save(
        [
            ["순번", "성명", "생년월일", "기본급", "처우개선비", "두루누리 국민연금", "국민연금"],
            [1, "홍길동", "1970-01-01", 2000000, 50000, 10000, 90000],
        ],
        tmp_path / "src.xlsx",
    )
    dst = _dest(tmp_path / "dst.xlsx", ["기본급", "처우개선비"], ["국민연금", "두루누리 국민연금"], [("1", "홍길동")])
    result = process("x", src, dst)
    rec = result.records[0]
    assert rec.allowances["처우개선비"] == 50000
    assert rec.deductions["두루누리 국민연금"] == 10000
    assert rec.deductions["국민연금"] == 90000


def test_parenthesized_continuation_row_is_summed(tmp_path: Path) -> None:
    """업체054: 사원009 아래 "(목)" 줄의 금액은 사원009 것이다."""
    src = _save(
        [
            # 실제 파일처럼 순번 열 머리글이 비어 있고, 동반 단어 대신 급여 항목명만 있다
            [None, "성명", "기본급", "주휴수당", "국민연금", "건강보험"],
            [1, "사원009", 1413840, 254490, 0, 66360],
            [None, "(목)", 51600, None, None, None],
            [2, "사원011", 650160, 117030, 36810, 29220],
        ],
        tmp_path / "src.xlsx",
    )
    dst = _dest(tmp_path / "dst.xlsx", ["기본급"], ["국민연금"], [("1", "사원009"), ("2", "사원011")])
    result = process("x", src, dst)
    by_name = {r.사원명: r for r in result.records}
    assert by_name["사원009"].allowances["기본급"] == 1465440
    assert "(목)" not in by_name


def test_insurance_not_summed_from_helper_columns(tmp_path: Path) -> None:
    """같은 4대보험 이름이 오른쪽에 또 나오면(계산용 보조표) 더하지 않는다."""
    src = _save(
        [
            ["순번", "성명", "생년월일", "기본급", "국민연금", "건강보험", "공제총액", "국민연금", "건강보험", "기본급"],
            [1, "홍길동", "1970-01-01", 2000000, 90000, 70000, 160000, 2000000, 1, 2000000],
        ],
        tmp_path / "src.xlsx",
    )
    dst = _dest(tmp_path / "dst.xlsx", ["기본급"], ["국민연금", "건강보험"], [("1", "홍길동")])
    result = process("x", src, dst)
    rec = result.records[0]
    assert rec.deductions["국민연금"] == 90000
    assert rec.deductions["건강보험"] == 70000
    assert rec.allowances["기본급"] == 2000000


def test_prefix_fallback_for_long_basic_pay_label(tmp_path: Path) -> None:
    src = _save(
        [
            ["No", "성명", "생년월일", "기본급(식대 미포함)", "장기근속공제"],
            [1, "홍길동", "1970-01-01", 1956880, 6730],
        ],
        tmp_path / "src.xlsx",
    )
    dst = _dest(tmp_path / "dst.xlsx", ["기본급"], ["장기근속수당공제"], [("1", "홍길동")])
    result = process("x", src, dst)
    rec = result.records[0]
    assert rec.allowances["기본급"] == 1956880
    assert rec.deductions["장기근속수당공제"] == 6730
    # 접두어로 붙인 매칭은 노란색 "확인 필요"로 표시된다
    assert "기본급" in result.ambiguous_dest_keys
