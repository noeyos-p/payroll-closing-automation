"""Synthetic workbooks and temporary copies only; never save the live list."""
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import pytest

from core import payroll_company_list as module
from core.payroll_company_list import (
    ClosureResult, CompanyListError, ConcurrentModificationError,
    DuplicateCompanyError, is_completed, normalize_business_number,
    read_company_list, update_company_list,
)


HEADERS = ["번호", "업체명", "관리담당자", "사업자번호", "귀속지급기준", "사업소득여부",
           "원천세신고여부", "급여입력여부", "신고서저장\n완료여부", "신고구분",
           "기타1", "기타2", "확인", "비고"]
LIVE_PATH = Path("C:/Users/<사용자>/Desktop/업체리스트(10월지급).xlsx")


def make_book(tmp_path, rows=None):
    path = tmp_path / "companies.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "업체"
    ws.append(HEADERS)
    for row in rows or [[1, "테스트업체", "담당", "012-34-56789", "기준", "X", "O", None,
                         "X", "종류", "유지K", "유지L", "확인", "기존 비고"]]:
        ws.append(row)
    for address in ("H2", "N2", "D2"):
        ws[address].font = Font(name="맑은 고딕", bold=True, color="FF112233")
        ws[address].fill = PatternFill("solid", fgColor="FFABCDEF")
        ws[address].alignment = Alignment(wrap_text=True)
    ws["D2"].number_format = "0000000000"
    ws["K3"] = "=SUM(A2:A2)"
    ws.column_dimensions["N"].width = 42
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "E2"
    other = wb.create_sheet("다른시트")
    other["A1"] = "=업체!A2+1"
    other["B2"] = "보존"
    other["B2"].font = Font(italic=True)
    other.merge_cells("B2:C3")
    other.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return path


def read_hn(path, row=2):
    wb = load_workbook(path)
    try:
        return wb["업체"].cell(row, 8).value, wb["업체"].cell(row, 14).value
    finally:
        wb.close()


@pytest.mark.parametrize("value", ["O", "o", " O ", "○", "◯", "⭕", "〇", "⭘"])
def test_standalone_complete(value):
    assert is_completed(value)


@pytest.mark.parametrize("value", [None, "", "X", "O,X", "O,O", "○,X", "OX", 0, True])
def test_not_all_complete(value):
    assert not is_completed(value)


@pytest.mark.parametrize("value", ["0123456789", "012-34-56789", 123456789, 123456789.0])
def test_leading_zero(value):
    assert normalize_business_number(value) == "0123456789"


@pytest.mark.parametrize("value", [True, -1, 1.25, float("nan"), "=A1", "12345678901", "12a"])
def test_invalid_number(value):
    with pytest.raises(CompanyListError):
        normalize_business_number(value)


def test_read_rows_and_numeric_business_number(tmp_path):
    path = make_book(tmp_path)
    wb = load_workbook(path)
    wb["업체"]["D2"] = 123456789
    wb.save(path)
    wb.close()
    row = read_company_list(path).rows[0]
    assert row.row_number == 2
    assert row.business_number == "0123456789"
    assert row.company_name == "테스트업체"
    assert row.manager == "담당"
    assert row.payment_basis == "기준"
    assert row.business_income == "X"
    assert row.withholding_report == "O"
    assert row.report_saved == "X"
    assert row.report_kind == "종류"
    assert row.confirmation == "확인"
    assert row.notes == "기존 비고"
    assert not row.completed


def test_duplicates_are_visible_and_update_refuses(tmp_path):
    path = make_book(tmp_path, [[1, "중복", None, "0123456789"],
                                [2, "중복", None, 123456789]])
    snapshot = read_company_list(path)
    assert snapshot.duplicate_business_numbers == {"0123456789": (2, 3)}
    assert snapshot.duplicate_names == {"중복": (2, 3)}
    before = path.read_bytes()
    with pytest.raises(DuplicateCompanyError):
        snapshot.find_company(company_name="중복")
    with pytest.raises(DuplicateCompanyError):
        update_company_list(snapshot, {"0123456789": [ClosureResult("급여", True)]})
    assert path.read_bytes() == before
    assert not list(tmp_path.glob("*.bak"))


def test_same_name_different_numbers_requires_number(tmp_path):
    snapshot = read_company_list(make_book(tmp_path, [[1, "중복", None, "0123456789"],
                                                    [2, "중복", None, "1123456789"]]))
    with pytest.raises(DuplicateCompanyError):
        snapshot.find_company(company_name="중복")
    assert snapshot.find_company(business_number="1123456789").row_number == 3
    with pytest.raises(CompanyListError):
        snapshot.find_company(business_number="1123456789", company_name="다름")


@pytest.mark.parametrize("closures,expected,label", [
    ([ClosureResult("급여", True)], "O", "급여: O"),
    ([ClosureResult("추급", True), ClosureResult("급여", True), ClosureResult("급+상", True)],
     "O", "급여: O; 급+상: O; 추급: O"),
    ([ClosureResult("추급", True), ClosureResult("급여", False, "확인 필요"), ClosureResult("급+상", True)],
     "X,O,O", "급여: X (확인 필요); 급+상: O; 추급: O"),
    ([ClosureResult("추급", False, "자료 없음"), ClosureResult("급여", True)],
     "O,X", "급여: O; 추급: X (자료 없음)"),
    ([ClosureResult("급+상", False, "차이"), ClosureResult("급여", False, "미확인")],
     None, "급여: X (미확인); 급+상: X (차이)"),
    ([ClosureResult("급여", False)], None, "급여: X (미마감)"),
    ([ClosureResult("추급", True), ClosureResult("추급", False, "두 번째 파일")],
     None, "추급: X (두 번째 파일)"),
])
def test_closure_aggregation_and_note_idempotence(tmp_path, closures, expected, label):
    path = make_book(tmp_path)
    result = update_company_list(read_company_list(path), {"0123456789": closures})
    assert result.changed_rows == (2,)
    h, n = read_hn(path)
    assert h == expected
    assert n == "기존 비고\n[급여마감] " + label
    before = path.read_bytes()
    repeat = update_company_list(read_company_list(path), {"0123456789": reversed(closures)})
    assert repeat.backup_path is None
    assert path.read_bytes() == before
    assert len(list(tmp_path.glob("*.bak"))) == 1


@pytest.mark.parametrize("completed", ["O", "○", "o"])
def test_skip_existing_complete_entire_row(tmp_path, completed):
    path = make_book(tmp_path)
    wb = load_workbook(path)
    wb["업체"]["H2"] = completed
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    result = update_company_list(read_company_list(path), {"0123456789": [ClosureResult("급여", False, "실패")]})
    assert result.skipped_rows == (2,)
    assert result.changed_rows == ()
    assert result.backup_path is None
    assert path.read_bytes() == before


def test_backup_reopen_styles_other_sheets_formulas_and_only_hn(tmp_path):
    path = make_book(tmp_path)
    before = path.read_bytes()
    old = load_workbook(path)
    result = update_company_list(read_company_list(path), {"0123456789": [ClosureResult("급여", True)]})
    assert result.backup_path.read_bytes() == before
    new = load_workbook(path)
    try:
        assert old.sheetnames == new.sheetnames
        for old_ws, new_ws in zip(old, new):
            assert old_ws.sheet_state == new_ws.sheet_state
            assert str(old_ws.merged_cells) == str(new_ws.merged_cells)
            for cells in old_ws:
                for cell in cells:
                    updated = new_ws[cell.coordinate]
                    assert cell._style == updated._style
                    if old_ws.title != "업체" or cell.coordinate not in ("H2", "N2"):
                        assert (cell.value, cell.data_type) == (updated.value, updated.data_type)
        assert new["업체"].column_dimensions["N"].width == 42
        assert new["업체"].row_dimensions[2].height == 30
    finally:
        old.close()
        new.close()


def test_partial_then_closed_and_history_preserved(tmp_path):
    path = make_book(tmp_path)
    update_company_list(read_company_list(path), {"0123456789": [ClosureResult("추급", False, "대기"), ClosureResult("급여", True)]})
    assert read_hn(path)[0] == "O,X"
    update_company_list(read_company_list(path), {"0123456789": [ClosureResult("추급", True), ClosureResult("급여", True)]})
    h, n = read_hn(path)
    assert h == "O"
    assert "추급: X (대기)" in n and "추급: O" in n


def test_external_change_refused(tmp_path):
    path = make_book(tmp_path)
    snapshot = read_company_list(path)
    wb = load_workbook(path)
    wb["업체"]["M2"] = "다른 편집"
    wb.save(path)
    wb.close()
    external = path.read_bytes()
    with pytest.raises(ConcurrentModificationError):
        update_company_list(snapshot, {"0123456789": [ClosureResult("급여", True)]})
    assert path.read_bytes() == external


def test_change_during_staging_refused(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    snapshot = read_company_list(path)
    external = snapshot._source + b"external editor"
    original_verify = module._verify

    def verify_and_change(data, expected):
        original_verify(data, expected)
        path.write_bytes(external)

    monkeypatch.setattr(module, "_verify", verify_and_change)
    with pytest.raises(ConcurrentModificationError):
        update_company_list(snapshot, {"0123456789": [ClosureResult("급여", True)]})
    assert path.read_bytes() == external
    assert not list(tmp_path.glob(".payroll-company-*"))


def test_failed_replace_keeps_original_and_backup(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    before = path.read_bytes()

    def locked(*args):
        raise PermissionError("locked")

    monkeypatch.setattr(module.os, "replace", locked)
    with pytest.raises(PermissionError):
        update_company_list(read_company_list(path), {"0123456789": [ClosureResult("급여", True)]})
    assert path.read_bytes() == before
    assert next(tmp_path.glob("*.bak")).read_bytes() == before
    assert not list(tmp_path.glob(".payroll-company-*"))


@pytest.mark.parametrize("closures", [[], [ClosureResult("기타", True)], [ClosureResult("급여", "O")]])
def test_invalid_results_do_not_write(tmp_path, closures):
    path = make_book(tmp_path)
    before = path.read_bytes()
    with pytest.raises(CompanyListError):
        update_company_list(read_company_list(path), {"0123456789": closures})
    assert path.read_bytes() == before


def test_formula_in_target_refused(tmp_path):
    path = make_book(tmp_path)
    wb = load_workbook(path)
    wb["업체"]["N2"] = "=1+1"
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    with pytest.raises(CompanyListError):
        update_company_list(read_company_list(path), {"0123456789": [ClosureResult("급여", True)]})
    assert path.read_bytes() == before


def test_real_workbook_copy_read_update_without_logging_entries(tmp_path):
    if not LIVE_PATH.exists():
        pytest.skip("desktop source not available")
    source_bytes = LIVE_PATH.read_bytes()
    copy = tmp_path / "desktop-copy.xlsx"
    shutil.copyfile(LIVE_PATH, copy)
    snapshot = read_company_list(copy)
    assert snapshot.rows
    eligible = [r for r in snapshot.rows if r.business_number and not r.completed
                and r.business_number not in snapshot.duplicate_business_numbers]
    if eligible:
        result = update_company_list(snapshot, {eligible[0].business_number: [ClosureResult("급여", False, "사본 검증")]})
        assert result.backup_path.read_bytes() == source_bytes
        assert read_company_list(copy)
    assert LIVE_PATH.read_bytes() == source_bytes
