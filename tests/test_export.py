from __future__ import annotations

from pathlib import Path
import zipfile

import openpyxl

from core.engine import process
from core.export import export_wehago, validate_wehago_xlsx
from core.profile import ClientProfile
from core.target_schema import ColumnDef, build_target_column_order, drop_all_zero_columns
from core.wehago_schema import WehagoRecord


def _profile() -> ClientProfile:
    return ClientProfile(
        id="test",
        client_name="테스트센터",
        version=1,
        created_at="2026-07-26T00:00:00+00:00",
        updated_at="2026-07-26T00:00:00+00:00",
        allowance_columns=[ColumnDef("기본급", "기본급", "allowance")],
        deduction_columns=[ColumnDef("국민연금", "국민연금", "deduction")],
    )


def test_export_uses_shared_strings_and_preserves_korean_data(tmp_path) -> None:
    output_path = tmp_path / "위하고-업로드.xlsx"
    records = [
        WehagoRecord(
            사원코드="1001",
            사원명="홍길동\x00",
            allowances={"기본급": 2_000_000},
            deductions={"국민연금": 90_000},
            지급액계=2_000_000,
            공제액계=90_000,
            차인지급액=1_910_000,
        )
    ]

    export_wehago(_profile(), records, output_path)
    validate_wehago_xlsx(output_path)

    with zipfile.ZipFile(output_path) as archive:
        assert "xl/sharedStrings.xml" in archive.namelist()
        assert "사원명" in archive.read("xl/sharedStrings.xml").decode("utf-8")
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert 't="inlineStr"' not in sheet_xml
        assert 't="s"' in sheet_xml

    workbook = openpyxl.load_workbook(output_path, data_only=False, read_only=True)
    worksheet = workbook["Sheet1"]
    assert [worksheet.cell(1, col).value for col in range(1, 10)] == [
        "사원코드",
        "사원명",
        "기본급",
        "비과세합계",
        "과세합계",
        "지급액계",
        "국민연금",
        "공제액계",
        "차인지급액",
    ]
    assert worksheet.cell(2, 2).value == "홍길동"
    assert worksheet.cell(2, 3).value == 2_000_000
    assert worksheet.cell(3, 1).value == "합계"
    assert worksheet.cell(3, 9).value == 1_910_000
    workbook.close()


def test_export_round_trips_processed_records(tmp_path) -> None:
    fixtures = Path(__file__).parent / "fixtures" / "golden"
    result = process(
        "업체075",
        fixtures / "sample1-source.xlsx",
        fixtures / "sample1-wehago-before.xlsx",
    )
    assert result is not None

    output_path = tmp_path / "업체030-위하고업로드.xlsx"
    export_wehago(result.profile, result.records, output_path)

    allowance_columns = drop_all_zero_columns(result.profile.allowance_columns, result.records, "allowance")
    deduction_columns = drop_all_zero_columns(result.profile.deduction_columns, result.records, "deduction")
    columns = build_target_column_order(allowance_columns, deduction_columns)

    workbook = openpyxl.load_workbook(output_path, data_only=False, read_only=True)
    worksheet = workbook["Sheet1"]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    assert len(rows) == len(result.records) + 2
    assert list(rows[0]) == [column.label for column in columns]
    assert [row[1] for row in rows[1:-1]] == [record.사원명 for record in result.records]

    net_col_idx = [column.key for column in columns].index("차인지급액")
    assert sum(row[net_col_idx] for row in rows[1:-1]) == sum(record.차인지급액 for record in result.records)
    assert rows[-1][net_col_idx] == sum(record.차인지급액 for record in result.records)
