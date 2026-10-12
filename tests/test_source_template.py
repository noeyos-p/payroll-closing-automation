"""사람이 손으로 다시 짠 급여대장 수정본에서 "원본 -> 수정본 양식" 규칙을 배우고
다음 달 원본에 적용하기(core/source_template.py, 2026-10-08)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from core.source_template import (
    LAST_APPLY_FAILURE,  # noqa: F401 — 모듈 속성 존재 확인
    _eval_formula,
    _header_signature,
    apply_recipe,
    learn_recipe,
)
import core.source_template as st

PEOPLE = [("김철수", 2000000, 100000, 90000, 30000), ("사원015", 1800000, 0, 81000, 20000), ("사원016", 2100000, 50000, 94500, 41000)]


def _original(path: Path, month: int, people, extra_col: bool = False) -> None:
    """원본 급여대장 — 제목/직원정보/근로시간이 섞인 사업장 자체 양식."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"2026년{month:02d}월 급여대장"
    ws["A1"] = f"2026년 {month:02d}월분 급여대장"
    head = ["번호", "성 명", "근로시간", "기본급", "연장수당", "국민연금", "소득세", "실지급액"]
    if not extra_col:
        head.remove("연장수당")
    ws.append([])
    ws.append(head)
    for i, (name, base, ot, pension, tax) in enumerate(people, start=1):
        row = [i, " ".join(name), 209, base]
        if extra_col:
            row.append(ot)
        row += [pension, tax, base + (ot if extra_col else 0) - pension - tax]
        ws.append(row)
    ws.append(["합계"])
    wb.save(path)


def _fixed(path: Path, people) -> None:
    """사람이 이지케어 표준양식(★이지케어양식(함수)★ver2)으로 다시 짠 수정본 —
    0원인 연장수당 열은 빼고, 근로시간도 뺐다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "임금대장"
    ws["A1"] = "2026년 08월 임금대장(상세)"
    ws["A3"], ws["B3"], ws["C3"], ws["E3"], ws["G3"], ws["H3"] = "No", "이름", "수당", "공제", "공제액계", "지급 급여"
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:D3")
    ws.merge_cells("E3:F3")
    ws["C4"], ws["E4"], ws["F4"] = "기본급", "국민연금", "소득세"
    for i, (name, base, _ot, pension, tax) in enumerate(people, start=1):
        r = 4 + i
        ws.cell(r, 1, i)
        ws.cell(r, 2, name)
        ws.cell(r, 3, base)
        ws.cell(r, 4, base)  # 수당 합계(수식 대신 값)
        ws.cell(r, 5, pension)
        ws.cell(r, 6, tax)
        ws.cell(r, 7, pension + tax)
        ws.cell(r, 8, base - pension - tax)
    wb.save(path)


def test_eval_formula_handles_sum_and_arithmetic() -> None:
    vals = {3: 100.0, 4: 50.0, 5: 25.0}
    assert _eval_formula("=SUM(C{r}:E{r})", 7, vals) == 175.0
    assert _eval_formula("=C{r}-D{r}", 7, vals) == 50.0
    assert _eval_formula("=C1+D{r}", 7, vals) is None  # 다른 행 참조는 계산하지 않는다


def test_header_signature_ignores_month_specific_titles(tmp_path: Path) -> None:
    for month in (8, 9):
        _original(tmp_path / f"o{month}.xlsx", month, PEOPLE)
    sigs = []
    for month in (8, 9):
        ws = openpyxl.load_workbook(tmp_path / f"o{month}.xlsx")[f"2026년{month:02d}월 급여대장"]
        sigs.append(_header_signature(ws, 4, 4))
    assert sigs[0] == sigs[1] == "기본급"


def test_learn_and_apply_same_month_reproduces_fixed(tmp_path: Path) -> None:
    orig, fixed = tmp_path / "가나다센터 8월귀속 9월지급 급여대장.xlsx", tmp_path / "fixed.xlsx"
    _original(orig, 8, PEOPLE)
    _fixed(fixed, PEOPLE)
    recipe = learn_recipe(orig, fixed)
    assert recipe is not None and recipe.complete, recipe.problems if recipe else None
    ws, layout, notes = apply_recipe(recipe, orig, fixed)
    names = [ws.cell(layout.data_start_row + i, 2).value for i in range(3)]
    assert names == ["김철수", "사원015", "사원016"]  # 원본의 "김 철 수" 띄어쓰기는 붙인다
    assert ws.cell(layout.data_start_row, 3).value == 2000000
    assert ws.cell(layout.data_start_row + 2, 6).value == 41000


def test_next_month_new_column_is_carried_over_and_sheet_rename_is_followed(tmp_path: Path) -> None:
    orig8, fixed = tmp_path / "가나다센터 8월귀속 9월지급 급여대장.xlsx", tmp_path / "fixed.xlsx"
    _original(orig8, 8, PEOPLE)
    _fixed(fixed, PEOPLE)
    recipe = learn_recipe(orig8, fixed)
    assert recipe.complete

    # 다음 달: 시트 이름이 바뀌고, 지난달엔 없던 "연장수당" 열이 중간에 생김 —
    # 열이 한 칸씩 밀려도 머리글로 다시 찾고, 새 열은 버리지 않고 덧붙인다.
    orig9 = tmp_path / "가나다센터 9월귀속 10월지급 급여대장.xlsx"
    _original(orig9, 9, PEOPLE, extra_col=True)
    out = apply_recipe(recipe, orig9, fixed)
    assert out is not None, st.LAST_APPLY_FAILURE
    ws, layout, notes = out
    assert any("연장수당" in n for n in notes)
    extra_col = next(c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value == "연장수당")
    assert ws.cell(3, extra_col).value == "수당"
    assert ws.cell(layout.data_start_row, extra_col).value == 100000
    assert ws.cell(layout.data_start_row, 6).value == 30000  # 소득세는 밀린 자리에서 다시 찾음


def test_wrong_column_is_caught_by_check_and_not_used(tmp_path: Path) -> None:
    orig8, fixed = tmp_path / "가나다센터 8월귀속 9월지급 급여대장.xlsx", tmp_path / "fixed.xlsx"
    _original(orig8, 8, PEOPLE)
    _fixed(fixed, PEOPLE)
    recipe = learn_recipe(orig8, fixed)

    # 다음 달 원본의 실지급액이 항목 합과 안 맞으면(열을 잘못 잡았거나 원본이
    # 바뀜) 변환하지 않고 None — 호출하는 쪽이 원래 방식으로 처리한다.
    orig9 = tmp_path / "가나다센터 9월귀속 10월지급 급여대장.xlsx"
    _original(orig9, 9, PEOPLE)
    wb = openpyxl.load_workbook(orig9)
    ws = wb.active
    for r in range(4, 7):
        ws.cell(r, 7, (ws.cell(r, 7).value or 0) + 12345)
    wb.save(orig9)
    assert apply_recipe(recipe, orig9, fixed) is None
    assert "검산" in st.LAST_APPLY_FAILURE
