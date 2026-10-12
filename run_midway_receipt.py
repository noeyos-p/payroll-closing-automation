"""중도퇴사자 연말정산 불러오기 → 귀속연도 확인 → 원천징수영수증 PDF (2026-10-12, 업체001에서 관측한 순서).

실행(작업 폴더 = 바탕화면\\　):
  "All2 - 복사본/.venv/Scripts/python.exe" "All2 - 급여입력 마감 자동화/run_midway_receipt.py" "<업체리스트 업체명>" [--search-name "<위하고 이름>"]

순서: 급여 화면 'A' 메뉴 → 연말정산 근로소득원천징수영수증(SWYM0101) → 중도 → 정산년월 지급월~지급월 →
조회 → 0명이면 F7 "중도 사원을 모두 불러오시겠습니까?" 확인 → 완료 안내 확인 → 사원별 귀속연도 끝(dt_rvrsto).
남길 사람(사용자 규칙): 차월·당월차월 = 귀속월(지급월−1) 퇴사, 당월 = 지급월 퇴사. 그 밖의 사람이 있으면
삭제는 아직 관측 전이라 멈춘다. 남길 사람만 체크 → #print → PDF로
작업파일\\Pdf\\{업체명} {귀속월}월귀속 중도퇴사자 원천징수영수증.pdf (인쇄형태 기본값 '_전체' 그대로, 사용자 결정).
연말정산 계산하기/계산완료/완료/마감/전사원불러오기/재계산/#delete 는 누르지 않는다.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

LIST = Path(r"C:\Users\<사용자>\Desktop\업체리스트(10월지급).xlsx")
PDF_DIR = Path(r"C:\Users\<사용자>\Desktop\작업파일\Pdf")
YEAR, PAY_MONTH = 2026, 10
SCREEN = "연말정산 근로소득원천징수영수증"
LOAD_PROMPT = "중도 사원을 모두 불러오시겠습니까?"
LOAD_DONE = "사원 불러오기가 완료 되었습니다."

CANVAS_JS = """() => [...document.querySelectorAll('canvas')].map(c => { const r = c.getBoundingClientRect();
    return {left: r.left, top: r.top, width: r.width, height: r.height}; }).filter(r => r.width > 100 && r.height > 100)"""
ROWS_JS = """() => { const g = window.Grids && window.Grids.getActiveGrid(); if (!g) return null;
    return g.getDataSource().getJsonRows().filter(r => r.cd_emp).map(r => ({cd_emp: r.cd_emp, nm_krname: r.nm_krname,
      dt_rvrsfrom: r.dt_rvrsfrom, dt_rvrsto: r.dt_rvrsto, ym_settle: r.ym_settle})); }"""
DIALOGS_JS = """() => [...document.querySelectorAll('.WSC_LUXDialog, .WSC_LUXDraggableDialog')]
    .filter(e => e.getBoundingClientRect().width > 50).map(e => e.innerText.trim().slice(0, 300))"""


class Stop(Exception):
    pass


def keep_month(period: str) -> int:
    """Retirement month that belongs to this report (user rule 2026-10-12)."""
    period = (period or "").replace("\n", "").strip()
    if period in ("차월", "당월차월"):
        return PAY_MONTH - 1
    if period == "당월":
        return PAY_MONTH
    raise Stop(f"업체리스트 귀속지급구분을 알 수 없음: {period!r}")


async def main(args, report: dict) -> None:
    from core import wehago_upload as wu
    from core.payroll_browser import PayrollBrowserError, WehagoPayrollAdapter
    from core.payroll_company_list import read_company_list
    from core.payroll_workflow import PaymentIdentity
    from core.print_pdf import PRINT_DIALOG_TITLE, find_window, print_dialog_contents, save_open_dialog_as_pdf, window_handles
    from core.settings import load_settings

    def log(*parts):
        print(" ".join(map(str, parts)), flush=True)

    rows = [r for r in read_company_list(LIST).rows if r.company_name == args.company]
    if len(rows) != 1 or not rows[0].business_number:
        raise Stop("업체리스트에서 업체를 하나로 찾지 못함")
    number = rows[0].business_number
    from openpyxl import load_workbook
    ws = load_workbook(LIST, read_only=True)["업체"]
    sheet_row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == args.company)
    if str(sheet_row[2] or "").strip() == "케어24":
        raise Stop("케어24 업체는 중도퇴사자 조회를 하지 않음(사용자 규칙 2026-10-12)")
    period = sheet_row[4]
    keep = keep_month(period)
    report.update(company=args.company, number=number, period=period, keep_month=keep)
    accrual = PAY_MONTH - 1
    dest = PDF_DIR / f"{args.company} {accrual}월귀속 중도퇴사자 원천징수영수증.pdf"
    if dest.exists():
        raise Stop(f"PDF가 이미 있음: {dest.name}")

    out = ROOT / "data" / "live_runs" / ("midway_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
    out.mkdir(parents=True, exist_ok=True)
    report["out"] = str(out)
    s = load_settings()
    main_page = await wu._get_main_page(300_000, s.wehago_id, s.wehago_password)
    for origin in ("https://www.wehago.com", "https://smarta.wehago.com"):
        await main_page.context.grant_permissions(["local-network-access"], origin=origin)  # print agent on 127.0.0.1
    adapter = WehagoPayrollAdapter(main_page, {number: args.company}, out / "dl", allow_mutation=False)
    opening = PaymentIdentity(number, YEAR, accrual, PAY_MONTH, None, "급여")
    try:
        page = await adapter._company_page(opening)
    except PayrollBrowserError as exc:
        if "검색 결과가 없습니다" not in str(exc) or not args.search_name:
            raise
        log("업체 검색어 재시도:", args.search_name)
        adapter.companies[number] = args.search_name
        page = await adapter._company_page(opening)

    # 'A' menu (top-left) → 연말정산관리 column → the receipt screen, same tab.
    link = page.get_by_text(SCREEN, exact=True)
    for _ in range(3):
        await page.mouse.click(22, 77)
        await page.wait_for_timeout(1500)
        if await link.count() == 1:
            break
    if await link.count() != 1:
        raise Stop("전체메뉴에서 연말정산 근로소득원천징수영수증을 찾지 못함")
    await link.click()
    await page.wait_for_timeout(4000)
    if "SWYM0101" not in page.url:
        raise Stop("SWYM0101 화면이 열리지 않음")
    await page.get_by_role("button", name="중도", exact=True).click()
    await page.wait_for_timeout(2000)

    async def query():
        spans = page.locator("#UFOPeriodMonthField1 .fakeinput")
        if await spans.count() != 2:
            raise Stop("정산년월 칸 두 개를 찾지 못함")
        for i in range(2):
            await spans.nth(i).click()
            await page.wait_for_timeout(400)
            await page.keyboard.type(f"{YEAR}{PAY_MONTH:02d}")
            await page.keyboard.press("Enter")
            await page.wait_for_timeout(800)
        texts = [t.strip() for t in await spans.all_inner_texts()]
        if texts != [f"{YEAR}.{PAY_MONTH:02d}"] * 2:
            raise Stop(f"정산년월 입력 확인 실패: {texts}")
        await page.get_by_role("button", name="조회", exact=True).click()
        await page.wait_for_timeout(3000)
        left = min(await page.evaluate(CANVAS_JS), key=lambda r: r["left"])
        await page.mouse.click(left["left"] + left["width"] / 2, left["top"] + 40)
        await page.wait_for_timeout(800)
        return left, await page.evaluate(ROWS_JS) or []

    left, people = await query()
    log("조회", [(p["nm_krname"], p["dt_rvrsto"]) for p in people])
    if not people:
        await page.keyboard.press("F7")
        await page.wait_for_timeout(2500)
        prompt = page.locator(".WSC_LUXDialog").filter(has_text=LOAD_PROMPT)
        if await prompt.count() != 1:
            raise Stop("F7 불러오기 확인창이 예상과 다름: " + json.dumps(await page.evaluate(DIALOGS_JS), ensure_ascii=False))
        await prompt.get_by_role("button", name="확인", exact=True).click()
        await page.wait_for_timeout(4000)
        done = page.locator(".WSC_LUXDialog").filter(has_text=LOAD_DONE)
        if await done.count() == 1:
            await done.get_by_role("button", name="확인", exact=True).click()
            await page.wait_for_timeout(1000)
        others = await page.evaluate(DIALOGS_JS)
        if others:
            raise Stop("불러오기 후 처음 보는 창: " + json.dumps(others, ensure_ascii=False))
        report["loaded"] = True
        left, people = await query()
        log("불러온 뒤", [(p["nm_krname"], p["dt_rvrsto"]) for p in people])
    report["people"] = people
    if not people:
        report["result"] = "no_midway"
        return
    bad = [p for p in people if len(str(p["dt_rvrsto"] or "")) != 8
           or int(str(p["dt_rvrsto"])[:4]) != YEAR or int(str(p["dt_rvrsto"])[4:6]) != keep]
    if bad:
        report["delete_needed"] = bad
        await page.screenshot(path=str(out / "delete_needed.png"))
        raise Stop(f"{keep}월 퇴사가 아닌 사람이 있음(삭제 미관측이라 멈춤): "
                   + ", ".join(f"{p['nm_krname']}({p['dt_rvrsto']})" for p in bad))

    await page.evaluate("() => window.Grids.getActiveGrid().checkAll(true)")
    await page.wait_for_timeout(500)
    checked = await page.evaluate("() => window.Grids.getActiveGrid().getCheckedRows()")
    if len(checked) != len(people):
        raise Stop(f"사원 체크 확인 실패: {checked}")
    before = await asyncio.to_thread(window_handles)
    await page.locator("#print").click()
    print_win = await asyncio.to_thread(find_window, PRINT_DIALOG_TITLE, None, 40, before)
    if print_win is None:
        raise Stop("인쇄창(Duzon - PrintDialog)이 뜨지 않음")
    form = await asyncio.to_thread(print_dialog_contents, print_win)
    log("인쇄형태", form)
    if form != "연말정산근로소득원천징수영수증_전체":
        print_win.close()
        raise Stop(f"인쇄형태가 기본값과 다름: {form}")
    saved = await asyncio.to_thread(save_open_dialog_as_pdf, print_win, dest)
    report["pdf"] = str(saved)
    report["result"] = "saved"
    log("PDF 저장", saved.name, saved.stat().st_size)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("company")
    parser.add_argument("--search-name")
    args = parser.parse_args()
    report: dict = {}
    try:
        asyncio.run(main(args, report))
    except Stop as exc:
        report.setdefault("result", "stopped")
        report["note"] = str(exc)
    except Exception as exc:
        report["result"] = "error"
        report["note"] = f"{type(exc).__name__}: {exc}"
    print("RESULT", json.dumps(report, ensure_ascii=False, default=str))
