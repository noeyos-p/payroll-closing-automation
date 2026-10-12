"""Annual cumulative ledgers (one sheet, 1월…12월 blocks stacked) → one month as
a plain one-row-per-person ledger the engine already reads.

Confirmed real case 업체046 (2026-10-12): '급여대장' sheet, column A
holds the month label on each block's first row, B the name, a per-block total row
(name blank or "12명"), and 류성희 listed *below* that total. There is no 실지급액
column; the person who entered September used 급여총액 − 공제총액 and merged a few
rows by hand. Those per-company decisions live in data/month_block_rules.json
(user decision 2026-10-12: "9월처럼 자동 처리") so nothing here guesses.
"""
from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from core.canon import norm
from core.xlsx_load import load_workbook as load_workbook_safe

RULES_PATH = Path(__file__).resolve().parent.parent / "data" / "month_block_rules.json"
_COUNT_LABEL = re.compile(r"^\d+명$")


class MonthBlockError(ValueError):
    pass


def load_rules(company_name: str, path: Path = RULES_PATH) -> dict | None:
    if not path.exists():
        return None
    rules = json.loads(path.read_text(encoding="utf-8"))
    return rules.get(company_name)


def _num(value) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise MonthBlockError(f"금액 칸에 숫자가 아닌 값: {value!r}")
    # Excel float noise such as 7019.999999999999 is shown as 7,020.
    return Decimal(repr(value)).quantize(Decimal(1))


def extract_month(source: Path, month: int, rule: dict, out_path: Path) -> tuple[Path, list[str]]:
    """Write `month`'s block as a 수당/공제/실지급액 ledger. Returns (path, notes)."""
    wb = load_workbook_safe(source, data_only=True)
    if rule["sheet"] not in wb.sheetnames:
        raise MonthBlockError(f"'{rule['sheet']}' 시트가 없습니다.")
    ws = wb[rule["sheet"]]
    header = rule["header_row"]
    allow = [(c, ws.cell(header, c).value) for c in rule["allowance_cols"]]
    deduct = [(c, ws.cell(header, c).value) for c in rule["deduction_cols"]]
    for c, label in allow + deduct + [(rule["gross_col"], None), (rule["deduction_total_col"], None)]:
        if label is not None and not isinstance(label, str):
            raise MonthBlockError(f"머리글이 예상과 다릅니다({c}열).")
    if norm(str(ws.cell(header, rule["gross_col"]).value)) != norm(rule["gross_label"]) or \
            norm(str(ws.cell(header, rule["deduction_total_col"]).value)) != norm(rule["deduction_total_label"]):
        raise MonthBlockError("급여총액/공제총액 머리글 위치가 규칙과 다릅니다.")

    label = f"{month}월"
    starts = [r for r in range(header + 1, ws.max_row + 1)
              if isinstance(ws.cell(r, rule["month_col"]).value, str)
              and ws.cell(r, rule["month_col"]).value.strip() == label]
    if len(starts) != 1:
        raise MonthBlockError(f"'{label}' 구간 시작을 하나로 찾지 못했습니다({len(starts)}건).")
    start = starts[0]
    end = next((r for r in range(start + 1, ws.max_row + 1)
                if isinstance(ws.cell(r, rule["month_col"]).value, str)
                and ws.cell(r, rule["month_col"]).value.strip()), ws.max_row + 1)

    people: list[dict] = []
    total_row = None
    for r in range(start, end):
        name = ws.cell(r, rule["name_col"]).value
        gross = ws.cell(r, rule["gross_col"]).value
        is_person = isinstance(name, str) and name.strip() and not _COUNT_LABEL.match(name.strip())
        if not is_person:
            if gross not in (None, "") and total_row is None:
                total_row = r
            continue
        if total_row is not None and not rule.get("include_after_total"):
            break
        people.append({
            "name": name.strip(), "row": r, "after_total": total_row is not None,
            "allow": [_num(ws.cell(r, c).value) for c, _ in allow],
            "deduct": [_num(ws.cell(r, c).value) for c, _ in deduct],
            "gross": _num(gross), "deduct_total": _num(ws.cell(r, rule["deduction_total_col"]).value),
        })
    if total_row is None:
        raise MonthBlockError(f"'{label}' 구간의 합계 줄을 찾지 못했습니다.")
    notes = []
    above = [p for p in people if not p["after_total"]]
    printed = _num(ws.cell(total_row, rule["gross_col"]).value)
    if sum((p["gross"] for p in above), Decimal(0)) != printed:
        raise MonthBlockError(f"'{label}' 사원별 급여총액 합이 합계 줄({printed:,})과 다릅니다.")
    for p in people:
        if sum(p["allow"], Decimal(0)) != p["gross"] or sum(p["deduct"], Decimal(0)) != p["deduct_total"]:
            raise MonthBlockError(f"{p['name']}: 항목 합이 급여총액/공제총액과 다릅니다({p['row']}행).")
    after = [p["name"] for p in people if p["after_total"]]
    if after:
        notes.append(f"합계 줄 아래 사원 포함(규칙): {', '.join(after)}")

    # Hand merges then renames, exactly as recorded in the rule.
    merged: dict[str, dict] = {}
    order: list[str] = []
    for p in people:
        target = rule.get("merge", {}).get(p["name"], p["name"])
        if target in merged:
            m = merged[target]
            m["allow"] = [a + b for a, b in zip(m["allow"], p["allow"])]
            m["deduct"] = [a + b for a, b in zip(m["deduct"], p["deduct"])]
            m["gross"] += p["gross"]
            m["deduct_total"] += p["deduct_total"]
            notes.append(f"{p['name']} → {target} 합산(규칙)")
        else:
            merged[target] = dict(p)
            order.append(target)
    for src, dst in rule.get("merge", {}).items():
        if dst not in merged and any(p["name"] == src for p in people):
            raise MonthBlockError(f"합산 대상 {dst}이(가) 이번 달에 없습니다.")
    rows = []
    for key in order:
        p = merged[key]
        new = rule.get("rename", {}).get(key, key)
        if new != key:
            notes.append(f"{key} → {new} 이름 정리(규칙)")
        rows.append((new, p))
    names = [n for n, _ in rows]
    if len(set(names)) != len(names):
        raise MonthBlockError("정리 후 같은 이름이 둘 이상입니다.")

    out = Workbook()
    o = out.active
    o.title = "임금대장"
    o.cell(1, 1, f"{month}월 급여대장(연간 시트에서 추출)")
    first_ded = 3 + len(allow) + 1
    o.cell(3, 1, "No"); o.cell(3, 2, "이름"); o.cell(3, 3, "수당"); o.cell(3, first_ded, "공제")
    for i, (_, lab) in enumerate(allow):
        o.cell(4, 3 + i, lab or "기타")
    o.cell(4, 3 + len(allow), "계")
    for i, (_, lab) in enumerate(deduct):
        o.cell(4, first_ded + i, "기타공제" if norm(lab or "기타") == "기타" else lab)
    tot_col = first_ded + len(deduct)
    o.cell(3, tot_col, "공제액계"); o.cell(3, tot_col + 1, "실지급액")
    for i, (name, p) in enumerate(rows):
        r = 5 + i
        o.cell(r, 1, i + 1); o.cell(r, 2, name)
        for j, v in enumerate(p["allow"]):
            if v:
                o.cell(r, 3 + j, int(v))
        o.cell(r, 3 + len(allow), int(p["gross"]))
        for j, v in enumerate(p["deduct"]):
            if v:
                o.cell(r, first_ded + j, int(v))
        o.cell(r, tot_col, int(p["deduct_total"]))
        o.cell(r, tot_col + 1, int(p["gross"] - p["deduct_total"]))
    last = 5 + len(rows)
    o.cell(last, 1, "합계")
    o.cell(last, tot_col + 1, int(sum((p["gross"] - p["deduct_total"] for _, p in rows), Decimal(0))))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out.save(out_path)
    return out_path, notes
